"""Tests des exports CSV / JSON / Excel (sans base de données ni réseau)."""

from __future__ import annotations

import csv
import json

from yt_veille.reports import EXPORT_COLUMNS, export_rows, rows_for_domain  # noqa: F401


def _sample_rows() -> list[dict]:
    """Deux lignes incluant des champs d'affichage propres à l'interface web."""
    return [
        {
            "rank": 1,
            "score_total": 0.9123,
            "video_id": "abc",
            "title": "Vidéo une",
            "channel": "Chaîne",
            "view_count": 131888,
            "like_count": 4200,
            "comment_count": 310,
            "engagement_pct": 3.65,
            "published_at": "2026-01-01T00:00:00+00:00",
            "age_days": 2.0,
            "duration_seconds": 754,
            "score_views": 0.8,
            "score_engagement": 0.7,
            "score_growth": 0.9,
            "score_relevance": 0.5,
            "score_recency": 0.6,
            "url": "https://www.youtube.com/watch?v=abc",
            # Champs d'affichage uniquement
            "thumbnail_url": "https://i.ytimg.com/vi/abc/mqdefault.jpg",
            "duration_hms": "12:34",
            "channel_url": "https://www.youtube.com/channel/UCabc",
        },
        {
            "rank": 2,
            "score_total": 0.5,
            "video_id": "def",
            "title": "Vidéo deux",
            "channel": "",
            "view_count": 10,
            "like_count": 0,
            "comment_count": 0,
            "engagement_pct": 0.0,
            "published_at": "2026-01-02T00:00:00+00:00",
            "age_days": 1.0,
            "duration_seconds": 30,
            "score_views": 0.1,
            "score_engagement": 0.0,
            "score_growth": 0.2,
            "score_relevance": 0.3,
            "score_recency": 0.4,
            "url": "https://www.youtube.com/watch?v=def",
            "thumbnail_url": "",
            "duration_hms": "0:30",
            "channel_url": "",
        },
    ]


def test_export_csv_ignore_les_champs_daffichage(tmp_path):
    path = export_rows(_sample_rows(), "csv", tmp_path / "out.csv")
    with path.open(encoding="utf-8") as fh:
        reader = csv.DictReader(fh)
        assert reader.fieldnames == EXPORT_COLUMNS
        rows = list(reader)
    assert len(rows) == 2
    assert rows[0]["video_id"] == "abc"
    assert "thumbnail_url" not in rows[0]


def test_export_json_ne_garde_que_les_colonnes(tmp_path):
    path = export_rows(_sample_rows(), "json", tmp_path / "out.json")
    data = json.loads(path.read_text(encoding="utf-8"))
    assert len(data) == 2
    assert set(data[0]) == set(EXPORT_COLUMNS)


def test_export_xlsx(tmp_path):
    path = export_rows(_sample_rows(), "xlsx", tmp_path / "out.xlsx")
    assert path.exists()

    from openpyxl import load_workbook

    ws = load_workbook(path).active
    header = [c.value for c in ws[1]]
    assert header == EXPORT_COLUMNS
    assert ws.max_row == 3  # en-tête + 2 lignes


def test_duree_formatee():
    from yt_veille.reports import _hms

    assert _hms(30) == "0:30"
    assert _hms(754) == "12:34"
    assert _hms(3723) == "1:02:03"
    assert _hms(None) == "0:00"
